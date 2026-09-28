import csv
import os

# Find the folder where this Python file is located
folder = os.path.dirname(os.path.abspath(__file__))

# Build the CSV file path
file_path = os.path.join(folder, "sales_data.csv")


# ==========================================
# READ AND CLEAN CSV
# ==========================================

cleaned_data = []

with open(file_path, "r", newline="", encoding="utf-8") as file:

    reader = csv.DictReader(file)

    for row in reader:

        # Check missing quantity
        if row["Quantity"].strip() == "":

            print(
                f"⚠ Missing quantity for {row['Customer']} - "
                f"{row['Product']}"
            )

            # Practice business rule
            row["Quantity"] = "1"

            print("  → Quantity automatically set to 1")

        # Convert values to numbers
        quantity = int(row["Quantity"])
        price = float(row["Price"])

        # Calculate total sale
        total_sale = quantity * price

        # Add calculated value
        row["Total Sale"] = total_sale

        cleaned_data.append(row)


# ==========================================
# CALCULATE TOTAL REVENUE
# ==========================================

total_revenue = 0

for row in cleaned_data:

    total_revenue += row["Total Sale"]


# ==========================================
# DISPLAY RESULTS
# ==========================================

print("\n========== SALES REPORT ==========\n")

for row in cleaned_data:

    print(
        row["Customer"],
        "|",
        row["Product"],
        "| Quantity:",
        row["Quantity"],
        "| Price:",
        row["Price"],
        "| Total:",
        row["Total Sale"]
    )


print("\n================================")
print("TOTAL REVENUE:", total_revenue)
print("================================")


# ==========================================
# SALES BY PRODUCT
# ==========================================

product_sales = {}

for row in cleaned_data:

    product = row["Product"]
    total_sale = row["Total Sale"]

    if product not in product_sales:
        product_sales[product] = 0

    product_sales[product] += total_sale


print("\n========== SALES BY PRODUCT ==========\n")

for product, revenue in product_sales.items():
    print(product, ":", revenue)


# ==========================================
# BEST-SELLING PRODUCT
# ==========================================

best_product = max(product_sales, key=product_sales.get)
best_revenue = product_sales[best_product]

print("\n========== TOP PRODUCT ==========\n")

print("Top Product:", best_product)
print("Revenue Generated:", best_revenue)



# ==========================================
# SALES BY CUSTOMER
# ==========================================

customer_sales = {}

for row in cleaned_data:

    customer = row["Customer"]
    total_sale = row["Total Sale"]

    if customer not in customer_sales:
        customer_sales[customer] = 0

    customer_sales[customer] += total_sale


print("\n========== SALES BY CUSTOMER ==========\n")

for customer, revenue in customer_sales.items():
    print(customer, ":", revenue)

# ==========================================
# CREATE EXCEL REPORT
# ==========================================

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference


# Create workbook
workbook = Workbook()

# Rename first sheet
sheet = workbook.active
sheet.title = "Sales Report"


# ==========================================
# TITLE
# ==========================================

sheet["A1"] = "Automated Sales Report"
sheet["A1"].font = Font(size=18, bold=True)

sheet.merge_cells("A1:E1")


# ==========================================
# SUMMARY
# ==========================================

sheet["A3"] = "Total Revenue"
sheet["B3"] = total_revenue

sheet["A4"] = "Top Product"
sheet["B4"] = best_product

sheet["A5"] = "Top Product Revenue"
sheet["B5"] = best_revenue


# ==========================================
# PRODUCT SALES TABLE
# ==========================================

sheet["A7"] = "Product"
sheet["B7"] = "Revenue"

for cell in sheet[7]:
    cell.font = Font(bold=True)


row_number = 8

for product, revenue in product_sales.items():

    sheet.cell(row=row_number, column=1, value=product)
    sheet.cell(row=row_number, column=2, value=revenue)

    row_number += 1


# ==========================================
# BAR CHART
# ==========================================

chart = BarChart()

chart.title = "Revenue by Product"
chart.y_axis.title = "Revenue"
chart.x_axis.title = "Product"

data = Reference(
    sheet,
    min_col=2,
    min_row=7,
    max_row=row_number - 1
)

categories = Reference(
    sheet,
    min_col=1,
    min_row=8,
    max_row=row_number - 1
)

chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)

sheet.add_chart(chart, "D7")


# ==========================================
# CUSTOMER SALES
# ==========================================

sheet["A15"] = "Customer"
sheet["B15"] = "Revenue"

sheet["A15"].font = Font(bold=True)
sheet["B15"].font = Font(bold=True)

row_number = 16

for customer, revenue in customer_sales.items():

    sheet.cell(row=row_number, column=1, value=customer)
    sheet.cell(row=row_number, column=2, value=revenue)

    row_number += 1


# ==========================================
# PROFESSIONAL FORMATTING
# ==========================================

from datetime import datetime


# Header formatting
header_fill = PatternFill(
    "solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)


# Format Product Sales header
for cell in sheet[7]:

    if cell.value is not None:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center"
        )


# Format Customer Sales header
for cell in sheet[15]:

    if cell.value is not None:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center"
        )


# Currency formatting
for row in range(8, row_number):

    sheet.cell(
        row=row,
        column=2
    ).number_format = '₹#,##0.00'


for row in range(16, row_number):

    sheet.cell(
        row=row,
        column=2
    ).number_format = '₹#,##0.00'


# Summary currency formatting
sheet["B3"].number_format = '₹#,##0.00'
sheet["B5"].number_format = '₹#,##0.00'


# Summary labels
for cell in ["A3", "A4", "A5"]:

    sheet[cell].font = Font(
        bold=True
    )


# Freeze panes
sheet.freeze_panes = "A7"


# Enable filters
sheet.auto_filter.ref = "A7:B12"


# Column widths
sheet.column_dimensions["A"].width = 28
sheet.column_dimensions["B"].width = 22


# Add report timestamp
sheet["A20"] = "Report Generated"
sheet["B20"] = datetime.now().strftime(
    "%d-%m-%Y %H:%M:%S"
)

sheet["A20"].font = Font(bold=True)


# Alignment
for row in sheet.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="center"
        )


# ==========================================
# SAVE FILE
# ==========================================

folder = os.path.dirname(os.path.abspath(__file__))

report_path = os.path.join(
    folder,
    "sales_report.xlsx"
)

workbook.save(report_path)

print("\n================================")
print("EXCEL REPORT CREATED SUCCESSFULLY")
print("File:", report_path)
print("================================")