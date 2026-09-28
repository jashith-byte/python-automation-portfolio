import os
import requests

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ==========================================
# FETCH DATA FROM API
# ==========================================

url = "https://jsonplaceholder.typicode.com/users"

# ==========================================
# API REQUEST WITH ERROR HANDLING
# ==========================================

try:

    response = requests.get(
        url,
        timeout=10
    )

    response.raise_for_status()

    print(
        "API connection successful!"
    )

    print(
        "Status Code:",
        response.status_code
    )

    data = response.json()


except requests.exceptions.Timeout:

    print(
        "❌ API request timed out."
    )

    exit()


except requests.exceptions.RequestException as error:

    print(
        "❌ API request failed:",
        error
    )

    exit()

print("\n========== API DATA ==========\n")

# ==========================================
# ORGANIZE API DATA
# ==========================================

users = []

for user in data:

    user_info = {
        "Name": user["name"],
        "Email": user["email"],
        "Company": user["company"]["name"],
        "City": user["address"]["city"]
    }

    users.append(user_info)


# ==========================================
# DISPLAY ORGANIZED DATA
# ==========================================

print("\n========== ORGANIZED DATA ==========\n")

for user in users:

    print(
        user["Name"],
        "|",
        user["Email"],
        "|",
        user["Company"],
        "|",
        user["City"]
    )

# ==========================================
# ANALYZE API DATA
# ==========================================

# Total users
total_users = len(users)


# Find unique companies
companies = set()

for user in users:
    companies.add(user["Company"])

total_companies = len(companies)


# Count users by city
city_counts = {}

for user in users:

    city = user["City"]

    if city not in city_counts:
        city_counts[city] = 0

    city_counts[city] += 1


# Find city with the most users
top_city = max(
    city_counts,
    key=city_counts.get
)


# ==========================================
# DISPLAY ANALYSIS
# ==========================================

print("\n========== BUSINESS ANALYSIS ==========\n")

print("Total Users:", total_users)
print("Total Companies:", total_companies)

print("\nUsers by City:")

for city, count in city_counts.items():
    print(city, ":", count)

print("\nLargest User Group:", top_city)

# ==========================================
# EXPORT API DATA TO EXCEL
# ==========================================

# Create workbook
workbook = Workbook()

sheet = workbook.active
sheet.title = "API Report"


# ==========================================
# TITLE
# ==========================================

sheet["A1"] = "API Business Data Report"
sheet["A1"].font = Font(
    size=18,
    bold=True
)

sheet.merge_cells("A1:D1")


# ==========================================
# HEADERS
# ==========================================

headers = [
    "Name",
    "Email",
    "Company",
    "City"
]

for column, header in enumerate(headers, start=1):

    cell = sheet.cell(
        row=3,
        column=column,
        value=header
    )

    cell.font = Font(bold=True)


# ==========================================
# WRITE API DATA
# ==========================================

row_number = 4

for user in users:

    sheet.cell(
        row=row_number,
        column=1,
        value=user["Name"]
    )

    sheet.cell(
        row=row_number,
        column=2,
        value=user["Email"]
    )

    sheet.cell(
        row=row_number,
        column=3,
        value=user["Company"]
    )

    sheet.cell(
        row=row_number,
        column=4,
        value=user["City"]
    )

    row_number += 1


# ==========================================
# SUMMARY
# ==========================================

summary_row = row_number + 2

sheet.cell(
    row=summary_row,
    column=1,
    value="Total Users"
)

sheet.cell(
    row=summary_row,
    column=2,
    value=total_users
)

sheet.cell(
    row=summary_row + 1,
    column=1,
    value="Total Companies"
)

sheet.cell(
    row=summary_row + 1,
    column=2,
    value=total_companies
)


# ==========================================
# FORMATTING
# ==========================================

header_fill = PatternFill(
    "solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

for cell in sheet[3]:

    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center"
    )


sheet.column_dimensions["A"].width = 25
sheet.column_dimensions["B"].width = 32
sheet.column_dimensions["C"].width = 25
sheet.column_dimensions["D"].width = 20


# Freeze headers
sheet.freeze_panes = "A4"


# ==========================================
# SAVE REPORT
# ==========================================

folder = os.path.dirname(
    os.path.abspath(__file__)
)

report_path = os.path.join(
    folder,
    "api_business_report.xlsx"
)

workbook.save(report_path)


print("\n================================")
print("EXCEL REPORT CREATED")
print("File:", report_path)
print("================================")